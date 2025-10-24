
import pytest
from fastapi.testclient import TestClient
from unittest.mock import patch, MagicMock
import io
from docx import Document

from main import app
from security import get_current_user

# Override user dependency for all tests in this file
async def override_get_current_user():
    return {"uid": "test_user_uid"}

@pytest.fixture
def client():
    """Create a TestClient instance.""" 
    app.dependency_overrides[get_current_user] = override_get_current_user
    return TestClient(app)


def test_upload_docx_success(client):
    """Test successful upload and parsing of a .docx file."""
    # Arrange
    document = Document()
    document.add_paragraph("Hello, DOCX world!")
    file_stream = io.BytesIO()
    document.save(file_stream)
    file_stream.seek(0)

    # Act
    response = client.post(
        "/api/upload",
        files={"file": ("test.docx", file_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )

    # Assert
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "test.docx"
    assert "Hello, DOCX world!" in data["text"]

@patch('routers.upload.PdfReader')
def test_upload_pdf_success(mock_pdf_reader, client):
    """Test successful upload and parsing of a .pdf file by mocking PdfReader."""
    # Arrange
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Hello, PDF world!"
    mock_pdf_reader.return_value.pages = [mock_page]
    
    pdf_content = b'%PDF-1.4...' # Dummy content
    file_stream = io.BytesIO(pdf_content)

    # Act
    response = client.post(
        "/api/upload",
        files={"file": ("test.pdf", file_stream, "application/pdf")}
    )

    # Assert
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "test.pdf"
    assert "Hello, PDF world!" in data["text"]

def test_upload_txt_success(client):
    """Test successful upload and parsing of a .txt file."""
    # Arrange
    txt_content = b"Hello, TXT world!"
    file_stream = io.BytesIO(txt_content)

    # Act
    response = client.post(
        "/api/upload",
        files={"file": ("test.txt", file_stream, "text/plain")}
    )

    # Assert
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "test.txt"
    assert "Hello, TXT world!" in data["text"]

def test_upload_unsupported_type(client):
    """Test uploading a file with an unsupported content type."""
    # Arrange
    content = b"some image data"
    file_stream = io.BytesIO(content)

    # Act
    response = client.post(
        "/api/upload",
        files={"file": ("test.png", file_stream, "image/png")}
    )

    # Assert
    assert response.status_code == 400
    assert "Unsupported file extension. Please upload .pdf, .docx, or .txt" in response.text

@patch('routers.upload.Document')
def test_upload_empty_file(mock_document, client):
    """Test uploading a file that results in no extracted text."""
    # Arrange
    mock_document.return_value.paragraphs = []
    file_stream = io.BytesIO()
    # We need to save something to the stream for it to be a valid file
    doc = Document()
    doc.save(file_stream)
    file_stream.seek(0)

    # Act
    response = client.post(
        "/api/upload",
        files={"file": ("empty.docx", file_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )

    # Assert
    assert response.status_code == 400
    assert "Could not extract any text from the document. The file might be empty or scanned." in response.text
